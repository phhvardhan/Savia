"""
BrandDocument -> Word. All styling comes from brand/jnj_tokens.py; every
document is checked by brand/validator.py before it is returned.

Layout: a J&J Red cover band across the text width (eyebrow, headline, subhead in white),
then sections on white: red section headings, black body copy, and records
as blocks (bold title line, aligned bold labels) divided by thin rules.
"""
import io

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Emu, Inches, Pt, RGBColor

from app.brand import fonts
from app.brand import jnj_tokens as T
from app.brand.adapter import to_brand_document
from app.brand.document import BrandDocument, Item
from app.brand.validator import style_name, validate_docx

PAGE_W, MARGIN = Inches(8.5), Inches(0.9)
TEXT_W = Inches(8.5 - 2 * 0.9)
BAND_PAD = Inches(0.35)  # inner padding of the red band
LABEL_W = Inches(1.9)  # label column for record fields
BULLET_W = Inches(0.2)


def _twips(length) -> int:
    return int(Emu(length).inches * 1440)


def resolved_styles() -> dict:
    """Each role's concrete font family, size, weight and colours. Shared with the preview."""
    return {
        role: {
            "font_family": fonts.family(spec["font"]),
            "size_pt": spec["size_pt"],
            "bold": spec["bold"],
            "color": T.hex_for(spec["color"]),
            "background": T.hex_for(spec["background"]),
        }
        for role, spec in T.ROLES.items()
    }


def _add_styles(doc) -> None:
    for role, s in resolved_styles().items():
        style = doc.styles.add_style(style_name(role), WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = s["font_family"]
        style.font.size = Pt(s["size_pt"])
        style.font.bold = s["bold"]
        style.font.italic = False
        style.font.color.rgb = RGBColor.from_string(s["color"])
        fmt = style.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(s["size_pt"] * 0.5)
        fmt.line_spacing = 1.15
    doc.styles[style_name("section_heading")].paragraph_format.space_before = Pt(22)
    doc.styles[style_name("section_heading")].paragraph_format.keep_with_next = True
    doc.styles[style_name("item_title")].paragraph_format.space_before = Pt(10)
    doc.styles[style_name("item_title")].paragraph_format.keep_with_next = True


def _write(p, text: str, bold: bool | None = None) -> None:
    lines = text.split("\n")
    run = p.add_run(lines[0])
    run.bold = bold
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def _paragraph(container, role: str, text: str):
    p = container.add_paragraph(style=style_name(role))
    _write(p, text)
    return p


def _hanging(p, width) -> None:
    p.paragraph_format.left_indent = width
    p.paragraph_format.first_line_indent = -width


def _table(doc, widths: list):
    """Borderless fixed-width table with no cell padding (renders the same in Word, Pages, Quick Look)."""
    table = doc.add_table(rows=1, cols=len(widths))
    table.autofit = False
    table._tbl.tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{sum(_twips(w) for w in widths)}" w:type="dxa"/>'))
    table._tbl.tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
    table._tbl.tblPr.append(parse_xml(
        f'<w:tblCellMar {nsdecls("w")}><w:left w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>'))
    for col, w in zip(table.columns, widths):
        col.width = w
    for cell, w in zip(table.rows[0].cells, widths):
        cell.width = w
    return table


def _cover(doc, document: BrandDocument) -> None:
    """Red band across the text width; padding comes from paragraph indents and spacing."""
    cell = _table(doc, [TEXT_W]).rows[0].cells[0]
    red = T.hex_for(T.ROLES["headline"]["background"])
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{red}"/>'))

    rows = [(r, t) for r, t in (("eyebrow", document.eyebrow), ("headline", document.headline),
                                ("subhead", document.subhead)) if t]
    for i, (role, text) in enumerate(rows):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = doc.styles[style_name(role)]
        p.paragraph_format.left_indent = p.paragraph_format.right_indent = BAND_PAD
        _write(p, text)
    cell.paragraphs[0].paragraph_format.space_before = Pt(30)
    cell.paragraphs[-1].paragraph_format.space_after = Pt(26)


def _rule(p) -> None:
    color = T.hex_for(T.RULE_COLOR)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="8" w:color="{color}"/></w:pBdr>'))


def _item(doc, item: Item, first: bool) -> None:
    if item.title:
        p = _paragraph(doc, "item_title", item.title)
        if not first:
            _rule(p)
    if not item.fields:
        return
    table = _table(doc, [LABEL_W, TEXT_W - LABEL_W])
    for i, field in enumerate(item.fields):
        label, value = (table.rows[0] if i == 0 else table.add_row()).cells
        for cell, text, bold in ((label, field.label, True), (value, field.value, None)):
            p = cell.paragraphs[0]
            p.style = doc.styles[style_name("body")]
            p.paragraph_format.space_after = Pt(4)
            if bold:  # label column: keep a gap before the value even when a viewer auto-sizes columns
                p.paragraph_format.right_indent = Inches(0.2)
            _write(p, text, bold)


def render(document: BrandDocument):
    """BrandDocument -> python-docx Document (not yet validated)."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.left_margin, sec.right_margin = PAGE_W, MARGIN, MARGIN
    sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.8)
    _add_styles(doc)
    body = resolved_styles()["body"]
    doc.styles["Normal"].font.name = body["font_family"]
    doc.styles["Normal"].font.color.rgb = RGBColor.from_string(body["color"])

    _cover(doc, document)
    for section in document.sections:
        _paragraph(doc, "section_heading", section.heading)
        bullets = len(section.body) > 1
        for text in section.body:
            p = doc.add_paragraph(style=style_name("body"))
            if bullets:
                _hanging(p, BULLET_W)
                p.add_run("•  ")
            _write(p, text)
        for i, item in enumerate(section.items):
            _item(doc, item, first=i == 0)

    footer = sec.footer.paragraphs[0]
    footer.style = doc.styles[style_name("footer")]
    footer.add_run(document.footer)
    return doc


def render_docx_bytes(document: BrandDocument) -> bytes:
    doc = render(document)
    validate_docx(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_artifact_docx(title: str, subtitle: str, content: dict) -> bytes:
    return render_docx_bytes(to_brand_document(title, subtitle, content))
