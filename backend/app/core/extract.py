import os
import wave
import time
import tempfile
import av
from pypdf import PdfReader
from docx import Document
from app.core.config import settings
from app.utils.logger import logger

AUDIO_VIDEO_EXT = {".mp4", ".mov", ".mp3", ".wav", ".m4a", ".avi", ".mkv"}
DOC_EXT = {".txt", ".docx", ".pdf"}

_whisper_model = None


def get_whisper_model():
    global _whisper_model
    if _whisper_model is None:
        from faster_whisper import WhisperModel
        model_path = str(settings.WHISPER_MODEL_PATH)
        try:
            logger.info(f"Loading local Whisper model from path '{model_path}' on CPU...")
            _whisper_model = WhisperModel(model_path, device="cpu", compute_type="int8")
        except Exception as e:
            logger.warning(f"Failed to load Whisper model from '{model_path}': {e}. Falling back to default 'tiny'...")
            _whisper_model = WhisperModel("tiny", device="cpu", compute_type="int8")
        logger.info("Whisper model ready.")
    return _whisper_model



def extract_audio_to_wav(file_path: str, target_wav_path: str, sample_rate: int = 16000) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".wav":
        try:
            with wave.open(file_path, "rb") as w:
                if w.getnchannels() == 1 and w.getframerate() == sample_rate:
                    return file_path
        except Exception:
            pass

    logger.info(f"Extracting audio from {file_path} to 16kHz mono WAV...")
    container = av.open(file_path)
    audio_streams = [s for s in container.streams if s.type == "audio"]
    if not audio_streams:
        raise ValueError(f"No audio stream found in file: {file_path}")

    audio_stream = audio_streams[0]
    resampler = av.AudioResampler(format="s16", layout="mono", rate=sample_rate)

    with wave.open(target_wav_path, "wb") as wav_file:
        wav_file.setnchannels(1)
        wav_file.setsampwidth(2)
        wav_file.setframerate(sample_rate)

        for frame in container.decode(audio_stream):
            resampled_frames = resampler.resample(frame)
            for resampled_frame in resampled_frames:
                wav_file.writeframes(resampled_frame.to_ndarray().tobytes())

    container.close()
    logger.info(f"Audio extracted successfully: {target_wav_path}")
    return target_wav_path


def transcribe_with_groq(wav_path: str, api_key: str) -> str:
    from groq import Groq
    client = Groq(api_key=api_key)
    file_size = os.path.getsize(wav_path)
    max_bytes = 24 * 1024 * 1024
    start = time.time()

    if file_size <= max_bytes:
        logger.info(f"Sending {file_size / (1024*1024):.2f} MB audio to Groq Whisper API...")
        with open(wav_path, "rb") as f:
            resp = client.audio.transcriptions.create(
                file=(os.path.basename(wav_path), f.read()),
                model="whisper-large-v3-turbo",
                response_format="json",
                temperature=0.0,
            )
        text = resp.text.strip()
        logger.info(f"Groq STT completed in {time.time() - start:.1f}s")
        return text

    logger.info("Audio file size exceeds 24MB limit. Chunking into 10-minute slices...")
    transcripts = []
    with wave.open(wav_path, "rb") as wav_in:
        n_channels = wav_in.getnchannels()
        sampwidth = wav_in.getsampwidth()
        framerate = wav_in.getframerate()
        chunk_frames = framerate * 60 * 10
        total_frames = wav_in.getnframes()
        chunk_idx = 0

        while wav_in.tell() < total_frames:
            frames = wav_in.readframes(chunk_frames)
            if not frames:
                break
            chunk_file = f"{wav_path}_chunk_{chunk_idx}.wav"
            try:
                with wave.open(chunk_file, "wb") as wav_out:
                    wav_out.setnchannels(n_channels)
                    wav_out.setsampwidth(sampwidth)
                    wav_out.setframerate(framerate)
                    wav_out.writeframes(frames)

                with open(chunk_file, "rb") as f:
                    resp = client.audio.transcriptions.create(
                        file=(os.path.basename(chunk_file), f.read()),
                        model="whisper-large-v3-turbo",
                        response_format="json",
                        temperature=0.0,
                    )
                if resp.text and resp.text.strip():
                    transcripts.append(resp.text.strip())
            finally:
                if os.path.exists(chunk_file):
                    try:
                        os.remove(chunk_file)
                    except Exception:
                        pass
            chunk_idx += 1

    return " ".join(transcripts)






def transcribe_local_whisper_segments(wav_path: str) -> list[dict]:
    model = get_whisper_model()
    logger.info(f"Transcribing {wav_path} with timestamps using local Whisper...")

    segments, info = model.transcribe(
        wav_path,
        beam_size=5,
        vad_filter=True,
        vad_parameters=dict(min_silence_duration_ms=500),
        condition_on_previous_text=False,
    )

    result = []
    for seg in segments:
        if seg.text and seg.text.strip():
            result.append({
                "start": seg.start,
                "end": seg.end,
                "text": seg.text.strip(),
            })
    return result


def transcribe_groq_whisper_segments(wav_path: str, api_key: str) -> list[dict]:
    from groq import Groq
    client = Groq(api_key=api_key)
    file_size = os.path.getsize(wav_path)
    max_bytes = 24 * 1024 * 1024

    if file_size <= max_bytes:
        logger.info(f"Fast AI STT: Sending {file_size / (1024*1024):.2f} MB audio to Groq Whisper API...")
        with open(wav_path, "rb") as f:
            resp = client.audio.transcriptions.create(
                file=(os.path.basename(wav_path), f.read()),
                model="whisper-large-v3-turbo",
                response_format="verbose_json",
                temperature=0.0,
            )
        segments = []
        raw_segs = getattr(resp, "segments", None) or []
        for s in raw_segs:
            start = s.get("start") if isinstance(s, dict) else getattr(s, "start", 0.0)
            end = s.get("end") if isinstance(s, dict) else getattr(s, "end", 0.0)
            text = s.get("text") if isinstance(s, dict) else getattr(s, "text", "")
            if text and text.strip():
                segments.append({
                    "start": float(start),
                    "end": float(end),
                    "text": text.strip(),
                })
        if not segments and resp.text and resp.text.strip():
            segments.append({"start": 0.0, "end": 10.0, "text": resp.text.strip()})
        return segments

    # For audio > 24MB, chunk into 10-minute slices with Groq for blazing speed
    logger.info(f"Audio size {file_size / (1024*1024):.2f} MB exceeds 24MB limit. Chunking into 10-minute slices for Groq...")
    all_segments = []
    with wave.open(wav_path, "rb") as wav_in:
        n_channels = wav_in.getnchannels()
        sampwidth = wav_in.getsampwidth()
        framerate = wav_in.getframerate()
        chunk_frames = framerate * 60 * 10
        total_frames = wav_in.getnframes()
        chunk_idx = 0
        time_offset = 0.0

        while wav_in.tell() < total_frames:
            frames = wav_in.readframes(chunk_frames)
            if not frames:
                break
            chunk_file = f"{wav_path}_chunk_{chunk_idx}.wav"
            try:
                with wave.open(chunk_file, "wb") as wav_out:
                    wav_out.setnchannels(n_channels)
                    wav_out.setsampwidth(sampwidth)
                    wav_out.setframerate(framerate)
                    wav_out.writeframes(frames)

                with open(chunk_file, "rb") as f:
                    resp = client.audio.transcriptions.create(
                        file=(os.path.basename(chunk_file), f.read()),
                        model="whisper-large-v3-turbo",
                        response_format="verbose_json",
                        temperature=0.0,
                    )
                raw_segs = getattr(resp, "segments", None) or []
                for s in raw_segs:
                    start = s.get("start") if isinstance(s, dict) else getattr(s, "start", 0.0)
                    end = s.get("end") if isinstance(s, dict) else getattr(s, "end", 0.0)
                    text = s.get("text") if isinstance(s, dict) else getattr(s, "text", "")
                    if text and text.strip():
                        all_segments.append({
                            "start": round(time_offset + float(start), 2),
                            "end": round(time_offset + float(end), 2),
                            "text": text.strip(),
                        })
            finally:
                if os.path.exists(chunk_file):
                    try:
                        os.remove(chunk_file)
                    except Exception:
                        pass
            chunk_idx += 1
            time_offset += 600.0

    if all_segments:
        return all_segments
    return transcribe_local_whisper_segments(wav_path)



def transcribe_whisper_segments(wav_path: str, api_key: str | None = None) -> list[dict]:
    key = api_key or settings.GROQ_API_KEY
    if key:
        try:
            segs = transcribe_groq_whisper_segments(wav_path, key)
            if segs:
                logger.info(f"Groq Fast AI STT completed successfully ({len(segs)} segments).")
                return segs
        except Exception as e:
            logger.warning(f"Groq Whisper API failed ({e}), falling back to local Whisper...")

    return transcribe_local_whisper_segments(wav_path)


def transcribe_local_whisper(wav_path: str) -> str:
    segments = transcribe_local_whisper_segments(wav_path)
    if segments:
        return " ".join(s["text"] for s in segments if s.get("text"))
    return "[Transcription: No speech recognized in audio/video file]"


def transcribe_audio_video(file_path: str, api_key: str | None = None) -> str:
    key = api_key or settings.GROQ_API_KEY

    temp_wav = os.path.join(tempfile.gettempdir(), f"extracted_{os.getpid()}_{int(time.time())}.wav")

    try:
        extracted_wav = extract_audio_to_wav(file_path, temp_wav)
        if settings.WHISPER_MODEL_PATH.exists():
            logger.info(f"Prioritizing local cloned Whisper model from '{settings.WHISPER_MODEL_PATH}'...")
            try:
                return transcribe_local_whisper(extracted_wav)
            except Exception as e:
                logger.warning(f"Local Whisper STT failed, checking cloud STT fallback: {e}")
                if key:
                    try:
                        return transcribe_with_groq(extracted_wav, key)
                    except Exception as ge:
                        logger.warning(f"Groq STT also failed: {ge}")

        if key:
            try:
                return transcribe_with_groq(extracted_wav, key)
            except Exception as e:
                logger.warning(f"Groq STT failed, falling back to local Whisper: {e}")

        try:
            return transcribe_local_whisper(extracted_wav)
        except Exception as e:
            logger.error(f"All STT options failed: {e}")
            return f"[Transcription: Audio extracted from {os.path.basename(file_path)}, but Speech-to-Text transcription failed ({e}).]"
    finally:
        if os.path.exists(temp_wav):
            try:
                os.remove(temp_wav)
            except Exception:
                pass


def extract_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_docx(file_path: str) -> str:
    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def extract_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)
    text_chunks = []
    for page in reader.pages:
        txt = page.extract_text()
        if txt and txt.strip():
            text_chunks.append(txt.strip())
    return "\n\n".join(text_chunks) if text_chunks else "[PDF Document: No extractable text found]"


def extract_doc(file_path: str) -> str:
    """Extract text from legacy .doc files using docx fallback or raw text parsing."""
    try:
        doc = Document(file_path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        with open(file_path, "rb") as f:
            content = f.read()
            import re
            words = re.findall(rb"[\x20-\x7E\t\r\n]{4,}", content)
            return "\n".join(w.decode("latin1", errors="ignore") for w in words)


def extract_text_from_file(file_path: str, api_key: str | None = None) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext in AUDIO_VIDEO_EXT:
        return transcribe_audio_video(file_path, api_key=api_key)
    if ext in {".txt", ".vtt", ".srt", ".md", ".csv", ".json"}:
        return extract_txt(file_path)
    if ext == ".docx":
        return extract_docx(file_path)
    if ext == ".doc":
        return extract_doc(file_path)
    if ext == ".pdf":
        return extract_pdf(file_path)
    raise ValueError(f"Unsupported file type: {ext}")

