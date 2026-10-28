import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _label(key: str) -> str:
    return str(key).replace("_", " ").strip().capitalize()


def _cell_text(value: Any) -> str:
    """Flatten a value into readable cell text (no raw JSON)."""
    if value is None or (isinstance(value, str) and value.strip().lower() in {"none", "null", "n/a"}):
        return ""
    if isinstance(value, dict):
        keys = {k.lower() for k in value}
        if {"given", "when", "then"} <= keys:
            return f"Given {value.get('given', '')}\nWhen {value.get('when', '')}\nThen {value.get('then', '')}"
        return "\n".join(f"{_label(k)}: {_cell_text(v)}" for k, v in value.items() if v not in (None, "", [], {}))
    if isinstance(value, list):
        return "\n".join(f"• {_cell_text(v)}" if not isinstance(v, dict) else _cell_text(v) for v in value)
    return str(value)


def _record_columns(rows: list[dict]) -> list[str]:
    cols: list[str] = []
    for row in rows:
        for k in row:
            if k not in cols:
                cols.append(k)
    return cols


def _is_records(value: Any) -> bool:
    return isinstance(value, list) and bool(value) and all(isinstance(v, dict) for v in value)


def _unique_sections(content: dict) -> list[tuple[str, Any]]:
    """Top-level sections, skipping alias keys that repeat an earlier section
    (MOM stores decisions under both "decisions" and "decisions_finalized")."""
    seen: list[Any] = []
    out = []
    for key, value in content.items():
        if value in seen:
            continue
        seen.append(value)
        out.append((key, value))
    return out


def artifact_has_records(content: dict) -> bool:
    return any(_is_records(v) for v in content.values())


def artifact_to_xlsx(title: str, subtitle: str, content: dict) -> bytes:
    """One sheet per list of records (stories, action items, test cases, RTM rows…)."""
    sections = [(k, v) for k, v in _unique_sections(content) if _is_records(v)]
    if not sections:
        raise ValueError("This output has no tables to export to Excel.")

    header_fill = PatternFill("solid", start_color="1E293B")
    zebra = PatternFill("solid", start_color="F8FAFC")
    side = Side(border_style="thin", color="CBD5E1")
    border = Border(left=side, right=side, top=side, bottom=side)
    priority_styles = {
        "high": ("FEE2E2", "991B1B"), "critical": ("FEE2E2", "991B1B"),
        "medium": ("FEF3C7", "92400E"), "low": ("E0F2FE", "075985"),
    }

    wb = Workbook()
    wb.remove(wb.active)
    for key, rows in sections:
        ws = wb.create_sheet(_label(key)[:31])  # Excel caps sheet names at 31 chars
        cols = _record_columns(rows)
        ws.cell(row=1, column=1, value=title).font = Font(size=16, bold=True, color="0F172A")
        ws.cell(row=2, column=1, value=f"{subtitle} · {len(rows)} rows").font = Font(size=10, italic=True, color="64748B")

        for c, col in enumerate(cols, start=1):
            cell = ws.cell(row=4, column=c, value=_label(col))
            cell.fill, cell.border = header_fill, border
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        for r, row in enumerate(rows, start=5):
            for c, col in enumerate(cols, start=1):
                text = _cell_text(row.get(col))
                cell = ws.cell(row=r, column=c, value=text)
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if (r - 5) % 2 == 0:
                    cell.fill = zebra
                if col.lower() in {"priority", "severity"} and text.lower() in priority_styles:
                    fill, color = priority_styles[text.lower()]
                    cell.fill = PatternFill("solid", start_color=fill)
                    cell.font = Font(bold=True, color=color)

        # Width from content, capped so long criteria wrap instead of running off.
        for c, col in enumerate(cols, start=1):
            longest = max([len(_label(col))] + [max((len(line) for line in _cell_text(row.get(col)).split("\n")), default=0) for row in rows])
            ws.column_dimensions[get_column_letter(c)].width = min(max(12, longest + 2), 60)
        ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
